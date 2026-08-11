#!/usr/bin/env python3
"""Generate the one-page team worksheet for the S08 brand dossier (editable DOCX)."""

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Mm, Pt, RGBColor

RED = RGBColor(0xC1, 0x28, 0x1B)
GREY = RGBColor(0x77, 0x77, 0x77)
DARK = RGBColor(0x22, 0x22, 0x22)
BODY = "Georgia"
LABEL = "Arial Narrow"


def shade(cell, hex_fill):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(el)


def spaced(run, twentieths=30):
    """Letter-spacing for label runs."""
    el = OxmlElement("w:spacing")
    el.set(qn("w:val"), str(twentieths))
    run._element.get_or_add_rPr().append(el)


def label(cell, text, size=6.5, color=GREY, space=30, bold=False):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text.upper())
    r.font.name, r.font.size, r.font.color.rgb, r.bold = LABEL, Pt(size), color, bold
    spaced(r, space)
    return p


def row_height(row, mm, rule=WD_ROW_HEIGHT_RULE.AT_LEAST):
    row.height_rule = rule
    row.height = Mm(mm)


doc = Document()
s = doc.sections[0]
s.orientation, s.page_width, s.page_height = WD_ORIENT.PORTRAIT, Mm(210), Mm(297)
s.top_margin = s.bottom_margin = Mm(11)
s.left_margin = s.right_margin = Mm(12)

style = doc.styles["Normal"]
style.font.name, style.font.size, style.font.color.rgb = BODY, Pt(9.5), DARK
style.paragraph_format.space_after = Pt(0)

# ---------- header ----------
p = doc.add_paragraph()
r = p.add_run("COMMERCIAL PHOTOGRAPHY  /  RMIT UNIVERSITY VIETNAM")
r.font.name, r.font.size, r.font.color.rgb = LABEL, Pt(7), GREY
spaced(r, 44)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
r = p.add_run("THE BRAND ")
r.font.name, r.font.size, r.bold = LABEL, Pt(23), True
r2 = p.add_run("DOSSIER")
r2.font.name, r2.font.size, r2.bold, r2.font.color.rgb = LABEL, Pt(23), True, RED

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(7)
r = p.add_run("SESSION 08  /  TREND AND BRAND ANALYSIS  /  PRE-CLASS TEAM RESEARCH")
r.font.name, r.font.size, r.font.color.rgb = LABEL, Pt(8), DARK
spaced(r, 20)
pbdr = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "single")
bottom.set(qn("w:sz"), "18")
bottom.set(qn("w:color"), "111111")
pbdr.append(bottom)
p._p.get_or_add_pPr().append(pbdr)

# ---------- team block ----------
t = doc.add_table(rows=1, cols=3)
t.style = "Table Grid"
widths = (Mm(80), Mm(63), Mm(43))
row_height(t.rows[0], 13)
for cell, w, txt in zip(t.rows[0].cells, widths, ("Team name", "Brand chosen", "Product category")):
    cell.width = w
    label(cell, txt)
    cell.add_paragraph()

doc.add_paragraph().paragraph_format.space_after = Pt(3)

# ---------- members ----------
m = doc.add_table(rows=6, cols=4)
m.style = "Table Grid"
mw = (Mm(15), Mm(83), Mm(45), Mm(43))
hdr = m.rows[0]
row_height(hdr, 5)
for cell, w, txt in zip(hdr.cells, mw, ("No.", "Member name", "Student number", "Parts led")):
    cell.width = w
    shade(cell, "F0EDE6")
    label(cell, txt, size=6.5, color=RGBColor(0x55, 0x55, 0x55), bold=True)
for i in range(1, 6):
    row_height(m.rows[i], 9)
    for cell, w in zip(m.rows[i].cells, mw):
        cell.width = w
    p = m.rows[i].cells[0].paragraphs[0]
    r = p.add_run(str(i))
    r.font.name, r.font.size, r.font.color.rgb = BODY, Pt(11), GREY

# ---------- instruction ----------
p = doc.add_paragraph()
p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(7), Pt(5)
r = p.add_run("HOW TO USE THIS   ")
r.font.name, r.font.size, r.bold, r.font.color.rgb = LABEL, Pt(7), True, RED
spaced(r, 30)
r = p.add_run(
    "Divide the seven searches between the team. Write your headline finding on the line, and put the "
    "lead member's initials in the last column. Attach your slide deck, roughly one page per part with "
    "the images, and bring both to class."
)
r.font.name, r.font.size = BODY, Pt(9)

# ---------- seven parts ----------
PARTS = [
    ("01", "Brand", "Five adjectives, the owned colour with its hex code, and one line placing the brand against its nearest competitor."),
    ("02", "Genre", "The genre ratio across the last twenty posts (still life / lifestyle / portrait / food / fashion), and what that ratio reveals."),
    ("03", "Concept", "One named campaign and its concept in a single sentence. The concept is never the product itself."),
    ("04", "Strategy", "One strategic move from the last twelve months (launch, new audience, premiumisation, collaboration) and the images that carried it."),
    ("05", "Platform", "One product followed across three surfaces. What changed: crop, aspect ratio, negative space, how much the image carries alone."),
    ("06", "Brief", "The brief you reconstructed: deliverables, platforms, tone, mandatories, and what was clearly forbidden."),
    ("07", "Story", "The story in one sentence, and its structure: origin, ritual, transformation, or spectacle."),
]

pt = doc.add_table(rows=1 + len(PARTS), cols=3)
pt.style = "Table Grid"
pw = (Mm(20), Mm(139), Mm(27))

hdr = pt.rows[0]
row_height(hdr, 5)
for cell, w, txt in zip(hdr.cells, pw, ("Part", "The search, and the finding you bring", "Lead")):
    cell.width = w
    shade(cell, "111111")
    label(cell, txt, size=6.5, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)

for i, (num, name, question) in enumerate(PARTS, start=1):
    row = pt.rows[i]
    row_height(row, 20)
    for cell, w in zip(row.cells, pw):
        cell.width = w

    c = row.cells[0]
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(num)
    r.font.name, r.font.size, r.bold, r.font.color.rgb = LABEL, Pt(11), True, RED
    p2 = c.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r = p2.add_run(name.upper())
    r.font.name, r.font.size, r.font.color.rgb = LABEL, Pt(7), DARK
    spaced(r, 20)

    c = row.cells[1]
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(question)
    r.font.name, r.font.size, r.font.color.rgb = BODY, Pt(8), RGBColor(0x66, 0x66, 0x66)
    for _ in range(2):
        line = c.add_paragraph()
        line.paragraph_format.space_after = Pt(4)
        pbdr = OxmlElement("w:pBdr")
        b = OxmlElement("w:bottom")
        b.set(qn("w:val"), "dotted")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:color"), "AAAAAA")
        pbdr.append(b)
        line._p.get_or_add_pPr().append(pbdr)

    c = row.cells[2]
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("INITIALS")
    r.font.name, r.font.size, r.font.color.rgb = LABEL, Pt(6.5), RGBColor(0xAA, 0xAA, 0xAA)
    spaced(r, 20)

# ---------- footer ----------
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
pbdr = OxmlElement("w:pBdr")
top = OxmlElement("w:top")
top.set(qn("w:val"), "single")
top.set(qn("w:sz"), "10")
top.set(qn("w:color"), "111111")
pbdr.append(top)
p._p.get_or_add_pPr().append(pbdr)

f = doc.add_table(rows=1, cols=2)
fw = (Mm(96), Mm(90))
left, right = f.rows[0].cells
for cell, w in zip((left, right), fw):
    cell.width = w

label(left, "Attach and bring", size=6.5, color=RED, bold=True)
for item in (
    "Dossier deck, about one page per part, images included",
    "Part 05 screenshots, three platforms side by side",
    "This sheet, completed and signed off by the team",
):
    p = left.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("☐  " + item)
    r.font.name, r.font.size = BODY, Pt(8.5)

label(right, "Opens your pitch", size=6.5, color=RED, bold=True)
p = right.add_paragraph()
p.paragraph_format.space_after = Pt(3)
r = p.add_run("The one finding that surprised the team most:")
r.font.name, r.font.size = BODY, Pt(8.5)
for _ in range(2):
    line = right.add_paragraph()
    line.paragraph_format.space_after = Pt(4)
    pbdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "dotted")
    b.set(qn("w:sz"), "4")
    b.set(qn("w:color"), "AAAAAA")
    pbdr.append(b)
    line._p.get_or_add_pPr().append(pbdr)

out = "Brand-Dossier-Team-Worksheet.docx"
doc.save(out)
print("saved", out)
